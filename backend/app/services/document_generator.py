import io
from typing import Optional, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from app.models.report import Report


class DocumentGenerator:
    """Exports structured reports to PDF, DOCX, Markdown, and LaTeX."""

    @staticmethod
    def to_markdown(report: Report) -> str:
        """Returns clean Markdown formatted string."""
        header = f"# {report.title}\n\n"
        meta = f"**Type**: {report.report_type.replace('_', ' ').title()}  \n"
        meta += f"**Generated**: {report.created_at.strftime('%Y-%m-%d %H:%M:%S UTC')}  \n\n---\n\n"
        return header + meta + report.content_markdown

    @staticmethod
    def to_latex(report: Report) -> str:
        """Generates LaTeX document source."""
        clean_title = report.title.replace("_", "\\_").replace("&", "\\&").replace("%", "\\%")
        tex = [
            "\\documentclass[11pt,a4paper]{article}",
            "\\usepackage[utf8]{inputenc}",
            "\\usepackage{geometry}",
            "\\geometry{margin=1in}",
            "\\usepackage{hyperref}",
            "\\usepackage{titlesec}",
            "\\title{" + clean_title + "}",
            "\\author{IdeaLens AI Research Platform}",
            "\\date{\\today}",
            "\\begin{document}",
            "\\maketitle",
            "\\begin{abstract}",
            f"Automated evaluation and research document for {clean_title}.",
            "\\end{abstract}",
            "\\section{Report Overview}",
            f"\\textbf{{Report Type:}} {report.report_type}\\\\",
            "\\section{Content}",
            report.content_markdown.replace("#", "").replace("_", "\\_"),
            "\\end{document}"
        ]
        return "\n".join(tex)

    @staticmethod
    def to_docx_buffer(report: Report) -> io.BytesIO:
        """Builds a formatted Word (.docx) document."""
        doc = Document()

        # Document Title
        title_p = doc.add_heading(report.title, level=0)
        title_p.runs[0].font.color.rgb = RGBColor(20, 71, 216)  # Brand blue

        # Metadata
        meta_p = doc.add_paragraph()
        meta_p.add_run(f"Report Type: {report.report_type.replace('_', ' ').title()}\n").bold = True
        meta_p.add_run(f"Generated at: {report.created_at.strftime('%Y-%m-%d %H:%M UTC')}\n")
        meta_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph("―" * 40)

        # Parse markdown lines into headings & paragraphs
        lines = report.content_markdown.split("\n")
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith("- ") or stripped.startswith("* "):
                doc.add_paragraph(stripped[2:], style='List Bullet')
            else:
                doc.add_paragraph(stripped)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def to_pdf_buffer(report: Report) -> io.BytesIO:
        """Builds a styled PDF document using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#1447d8'),
            spaceAfter=12
        )
        
        meta_style = ParagraphStyle(
            'CustomMeta',
            parent=styles['Normal'],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#64748b'),
            spaceAfter=18
        )

        h2_style = ParagraphStyle(
            'CustomH2',
            parent=styles['Heading2'],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#0f266e'),
            spaceBefore=14,
            spaceAfter=6
        )

        body_style = ParagraphStyle(
            'CustomBody',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1e293b'),
            spaceAfter=8
        )

        bullet_style = ParagraphStyle(
            'CustomBullet',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1e293b'),
            leftIndent=15,
            spaceAfter=4
        )

        story = []
        # Title
        story.append(Paragraph(report.title, title_style))
        story.append(Paragraph(f"IdeaLens Innovation Report • {report.report_type.replace('_', ' ').title()} • {report.created_at.strftime('%Y-%m-%d')}", meta_style))
        story.append(Spacer(1, 10))

        # Parse content
        for line in report.content_markdown.split("\n"):
            s = line.strip()
            if not s:
                continue
            if s.startswith("#"):
                clean_heading = s.lstrip("#").strip()
                story.append(Paragraph(clean_heading, h2_style))
            elif s.startswith("- ") or s.startswith("* "):
                bullet_text = s[2:].strip()
                story.append(Paragraph(f"&bull; {bullet_text}", bullet_style))
            else:
                story.append(Paragraph(s, body_style))

        doc.build(story)
        buffer.seek(0)
        return buffer
