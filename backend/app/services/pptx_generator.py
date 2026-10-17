import io
from pptx import Presentation as PPTXPresentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from typing import List, Dict, Any, Optional
from app.models.presentation import Presentation, Slide


THEME_COLORS = {
    "navy_modern": {
        "bg": RGBColor(7, 13, 30),         # Dark Navy
        "card_bg": RGBColor(15, 25, 55),
        "text": RGBColor(241, 245, 249),
        "accent": RGBColor(47, 121, 246),   # Brand blue
        "subtext": RGBColor(148, 163, 184)
    },
    "minimalist_clean": {
        "bg": RGBColor(255, 255, 255),
        "card_bg": RGBColor(248, 250, 252),
        "text": RGBColor(15, 23, 42),
        "accent": RGBColor(30, 41, 59),
        "subtext": RGBColor(100, 116, 139)
    },
    "academic_slate": {
        "bg": RGBColor(241, 245, 249),
        "card_bg": RGBColor(255, 255, 255),
        "text": RGBColor(30, 41, 59),
        "accent": RGBColor(15, 118, 110),   # Teal
        "subtext": RGBColor(71, 85, 105)
    },
    "startup_cyber": {
        "bg": RGBColor(10, 15, 29),
        "card_bg": RGBColor(20, 30, 50),
        "text": RGBColor(248, 250, 252),
        "accent": RGBColor(139, 92, 246),   # Indigo / Purple
        "subtext": RGBColor(167, 139, 250)
    }
}


class PPTXGenerator:
    """Generates editable PowerPoint presentations using python-pptx."""

    @classmethod
    def generate_presentation_buffer(cls, presentation: Presentation, slides: List[Slide]) -> io.BytesIO:
        prs = PPTXPresentation()
        # Set 16:9 widescreen dimensions
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_slide_layout = prs.slide_layouts[6]  # Blank layout

        theme = THEME_COLORS.get(presentation.theme, THEME_COLORS["navy_modern"])

        for slide in sorted(slides, key=lambda s: s.sort_order):
            s_obj = prs.slides.add_slide(blank_slide_layout)

            # Slide background shape
            bg_shape = s_obj.shapes.add_shape(
                1,  # msoShapeRectangle
                Inches(0), Inches(0), prs.slide_width, prs.slide_height
            )
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = theme["bg"]
            bg_shape.line.fill.background()

            # Slide Title & Header Box
            title_box = s_obj.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
            tf = title_box.text_frame
            tf.word_wrap = True
            
            p = tf.paragraphs[0]
            p.text = slide.title
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = theme["accent"]

            if slide.subtitle:
                p_sub = tf.add_paragraph()
                p_sub.text = slide.subtitle
                p_sub.font.size = Pt(16)
                p_sub.font.color.rgb = theme["subtext"]

            # Content layout handling
            content_json = slide.content_json or {}
            bullets = content_json.get("bullets", [])

            if slide.layout == "title":
                # Large title layout
                p.font.size = Pt(44)
                p.alignment = PP_ALIGN.CENTER
                if slide.subtitle:
                    p_sub.font.size = Pt(22)
                    p_sub.alignment = PP_ALIGN.CENTER
                
                # Center author/meta note if present
                notes = content_json.get("notes") or f"IdeaLens Innovation Workspace • Theme: {presentation.theme.replace('_', ' ').title()}"
                center_box = s_obj.shapes.add_textbox(Inches(1.5), Inches(4.0), Inches(10.3), Inches(1.5))
                c_tf = center_box.text_frame
                c_tf.word_wrap = True
                c_p = c_tf.paragraphs[0]
                c_p.text = notes
                c_p.font.size = Pt(16)
                c_p.font.color.rgb = theme["subtext"]
                c_p.alignment = PP_ALIGN.CENTER

            elif slide.layout == "two_column":
                left_col = content_json.get("left_column", [])
                right_col = content_json.get("right_column", [])

                # Left Box
                left_box = s_obj.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.5))
                l_tf = left_box.text_frame
                l_tf.word_wrap = True
                for i, item in enumerate(left_col):
                    p_item = l_tf.paragraphs[0] if i == 0 else l_tf.add_paragraph()
                    p_item.text = f"• {item}"
                    p_item.font.size = Pt(18)
                    p_item.font.color.rgb = theme["text"]
                    p_item.space_after = Pt(12)

                # Right Box
                right_box = s_obj.shapes.add_textbox(Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.5))
                r_tf = right_box.text_frame
                r_tf.word_wrap = True
                for i, item in enumerate(right_col):
                    p_item = r_tf.paragraphs[0] if i == 0 else r_tf.add_paragraph()
                    p_item.text = f"• {item}"
                    p_item.font.size = Pt(18)
                    p_item.font.color.rgb = theme["text"]
                    p_item.space_after = Pt(12)

            else:
                # Default bullet layout
                content_box = s_obj.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
                c_tf = content_box.text_frame
                c_tf.word_wrap = True

                for i, b in enumerate(bullets):
                    p_b = c_tf.paragraphs[0] if i == 0 else c_tf.add_paragraph()
                    p_b.text = f"• {b}"
                    p_b.font.size = Pt(20)
                    p_b.font.color.rgb = theme["text"]
                    p_b.space_after = Pt(14)

                # Add citations / footer note if present
                citations = content_json.get("citations", [])
                if citations:
                    p_cite = c_tf.add_paragraph()
                    p_cite.text = f"Source References: {', '.join(citations)}"
                    p_cite.font.size = Pt(12)
                    p_cite.font.italic = True
                    p_cite.font.color.rgb = theme["subtext"]
                    p_cite.space_before = Pt(20)

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer
