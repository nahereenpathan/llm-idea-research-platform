import io
import os
import csv
import json
import uuid
import requests
import docx
from pptx import Presentation

BASE_URL = "http://127.0.0.1:8000"
ARTIFACTS_DIR = os.path.join(os.path.dirname(__file__), "test_downloaded_artifacts")
os.makedirs(ARTIFACTS_DIR, exist_ok=True)

def verify_all():
    print("=" * 70)
    print("PHASE 7 FINAL END-TO-END VERIFICATION SUITE")
    print("=" * 70)

    # -------------------------------------------------------------------------
    # 0. Health probe
    # -------------------------------------------------------------------------
    res = requests.get(f"{BASE_URL}/health")
    assert res.status_code == 200, f"Health probe failed: {res.text}"
    health_data = res.json()
    assert health_data["status"] == "healthy"
    assert health_data["database_connected"] is True
    print(f"[PASS] 0. Backend Health: OK ({health_data['platform']} v{health_data['version']})")

    # -------------------------------------------------------------------------
    # 1. Authenticate Seed User for Existing Saved Project 1
    # -------------------------------------------------------------------------
    login_res = requests.post(f"{BASE_URL}/api/v1/auth/login", json={
        "email": "researcher_mvp@idealens.ai",
        "password": "SecurePassword123!"
    })
    assert login_res.status_code == 200, f"Seed user login failed: {login_res.text}"
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("[PASS] 1. Authenticated as existing user: researcher_mvp@idealens.ai")

    # Verify existing saved Project 1
    p1_res = requests.get(f"{BASE_URL}/api/v1/projects/1", headers=headers)
    assert p1_res.status_code == 200, f"Project 1 fetch failed: {p1_res.text}"
    p1 = p1_res.json()
    project_id = p1["id"]
    assert p1["title"] == "Autonomous Drone Wildfire Early Detection"
    print(f"[PASS] 1b. Existing Saved Project 1 Retrieved: '{p1['title']}' (Domain: {p1['domain']})")

    # Check ideas in Project 1
    ideas_res = requests.get(f"{BASE_URL}/api/v1/projects/{project_id}/ideas", headers=headers)
    assert ideas_res.status_code == 200
    ideas = ideas_res.json()
    assert len(ideas) >= 1, "Project 1 should have saved ideas"
    selected_idea = next((i for i in ideas if i["is_selected"]), ideas[0])
    idea_id = selected_idea["id"]
    print(f"[PASS] 1c. Selected Candidate Idea: ID {idea_id} - '{selected_idea['title']}'")

    # -------------------------------------------------------------------------
    # 2. Real Document Generation for Existing Project 1
    # -------------------------------------------------------------------------
    print("\n--- Generating Publication-Grade Research Report ---")
    rep_res = requests.post(f"{BASE_URL}/api/v1/projects/{project_id}/reports/generate", json={
        "idea_id": idea_id,
        "report_type": "full_documentation",
        "title": "Autonomous Drone Wildfire Detection — Comprehensive Technical Dossier"
    }, headers=headers)
    assert rep_res.status_code in [200, 201], f"Report generation failed: {rep_res.text}"
    report = rep_res.json()
    report_id = report["id"]
    content_md = report["content_markdown"]
    print(f"[PASS] 2a. Report Generated (ID: {report_id}, Size: {len(content_md)} chars)")

    # Validate 10 mandatory architecture sections
    expected_sections = [
        "## 1. Executive Summary & Problem Space",
        "## 2. Idea Evaluation & Explainable Viability Analysis",
        "## 3. Academic Literature Review & Prior Art Evidence",
        "## 4. Patent Landscape & Intellectual Property Scoping",
        "## 5. Research Gaps & Validation Opportunities",
        "## 6. Technical Feasibility & Recommended Stack Architecture",
        "## 7. Implementation Roadmap & Phased Milestones",
        "## 8. Security, Privacy, Ethical & Business Risks",
        "## 9. Strategic Recommendations & Next Steps",
        "## 10. References & Source Citations"
    ]
    for s in expected_sections:
        assert s in content_md, f"Missing report section: {s}"
    print("[PASS] 2b. All 10 publication sections verified in generated Markdown")

    # -------------------------------------------------------------------------
    # 3. Content Integrity & Provenance Tagging Verification
    # -------------------------------------------------------------------------
    assert "[USER-PROVIDED INPUTS]" in content_md, "Missing [USER-PROVIDED INPUTS] tag"
    assert "[AI-DERIVED ANALYSIS]" in content_md, "Missing [AI-DERIVED ANALYSIS] tag"
    assert selected_idea["title"] in content_md, "Selected idea title missing from report"
    assert selected_idea["problem_statement"] in content_md, "Problem statement missing from report"
    print("[PASS] 3a. Verified provenance tags and user inputs integrated accurately")

    # Because Project 1 had 0 saved papers initially, verify that evidence limitation notice is present
    assert "EVIDENCE LIMITATION NOTICE" in content_md or "No academic papers" in content_md
    print("[PASS] 3b. Missing evidence correctly flagged with limitation notices (No fabrication)")

    # -------------------------------------------------------------------------
    # 4. Generate 10-Slide PPTX Presentation Deck
    # -------------------------------------------------------------------------
    print("\n--- Generating 10-Slide PowerPoint Presentation Deck ---")
    pres_res = requests.post(f"{BASE_URL}/api/v1/projects/{project_id}/presentations/generate", json={
        "idea_id": idea_id,
        "title": "Autonomous Drone Wildfire Detection — Executive Briefing",
        "theme": "navy_modern",
        "template_type": "startup_pitch"
    }, headers=headers)
    assert pres_res.status_code in [200, 201], f"Presentation generation failed: {pres_res.text}"
    pres = pres_res.json()
    pres_id = pres["id"]
    slides = pres["slides"]
    assert len(slides) == 10, f"Expected 10 slides, got {len(slides)}"
    print(f"[PASS] 4a. 10-Slide Presentation Deck Generated (ID: {pres_id})")

    # -------------------------------------------------------------------------
    # 5. File Downloads & Deep Binary Format Inspection
    # -------------------------------------------------------------------------
    print("\n--- Downloading & Inspecting Export Binaries on Disk ---")

    # 5A. PPTX Export
    pptx_url = f"{BASE_URL}/api/v1/presentations/{pres_id}/export?format=pptx"
    pptx_dl = requests.get(pptx_url, headers=headers)
    assert pptx_dl.status_code == 200, f"PPTX export HTTP status {pptx_dl.status_code}"
    assert pptx_dl.headers.get("content-type") == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert "attachment;" in pptx_dl.headers.get("content-disposition", "")
    assert ".pptx" in pptx_dl.headers.get("content-disposition", "")
    pptx_path = os.path.join(ARTIFACTS_DIR, "wildfire_pitch.pptx")
    with open(pptx_path, "wb") as f:
        f.write(pptx_dl.content)
    
    # Verify with python-pptx
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 10, f"Expected 10 slides, got {len(prs.slides)}"
    aspect_ratio = prs.slide_width / prs.slide_height
    assert abs(aspect_ratio - (16.0 / 9.0)) < 0.01, f"Expected 16:9 widescreen ratio, got {aspect_ratio}"
    titles = []
    for s in prs.slides:
        shape_texts = [p.text for shp in s.shapes if shp.has_text_frame for p in shp.text_frame.paragraphs if p.text]
        titles.append(shape_texts[0] if shape_texts else "Empty")
    print(f"[PASS] 5a. PPTX verified ({os.path.getsize(pptx_path)} bytes, 10 slides, 16:9 widescreen)")
    print(f"       Extracted Slide Titles: {titles[:4]} ... {titles[-2:]}")

    # 5B. PDF Export
    pdf_url = f"{BASE_URL}/api/v1/reports/{report_id}/export?format=pdf"
    pdf_dl = requests.get(pdf_url, headers=headers)
    assert pdf_dl.status_code == 200, f"PDF export HTTP status {pdf_dl.status_code}"
    assert pdf_dl.headers.get("content-type") == "application/pdf"
    assert "attachment;" in pdf_dl.headers.get("content-disposition", "")
    assert ".pdf" in pdf_dl.headers.get("content-disposition", "")
    assert pdf_dl.content.startswith(b"%PDF-"), "Invalid PDF header bytes"
    assert b"%%EOF" in pdf_dl.content[-1024:], "Invalid PDF trailer EOF marker"
    assert len(pdf_dl.content) > 3000, "PDF file suspiciously small"
    pdf_path = os.path.join(ARTIFACTS_DIR, "wildfire_report.pdf")
    with open(pdf_path, "wb") as f:
        f.write(pdf_dl.content)
    print(f"[PASS] 5b. PDF verified ({len(pdf_dl.content)} bytes, valid %PDF- header & %%EOF trailer)")

    # 5C. Word DOCX Export
    docx_url = f"{BASE_URL}/api/v1/reports/{report_id}/export?format=docx"
    docx_dl = requests.get(docx_url, headers=headers)
    assert docx_dl.status_code == 200, f"DOCX export HTTP status {docx_dl.status_code}"
    assert docx_dl.headers.get("content-type") == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "attachment;" in docx_dl.headers.get("content-disposition", "")
    assert ".docx" in docx_dl.headers.get("content-disposition", "")
    docx_path = os.path.join(ARTIFACTS_DIR, "wildfire_report.docx")
    with open(docx_path, "wb") as f:
        f.write(docx_dl.content)
    
    doc = docx.Document(docx_path)
    assert len(doc.paragraphs) > 25, f"DOCX has too few paragraphs: {len(doc.paragraphs)}"
    doc_full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "Autonomous Drone Wildfire Detection" in doc_full_text
    assert len(doc.tables) >= 1, "Expected table for metrics or roadmap in DOCX"
    print(f"[PASS] 5c. DOCX verified ({os.path.getsize(docx_path)} bytes, {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables)")

    # 5D. Markdown Export
    md_url = f"{BASE_URL}/api/v1/reports/{report_id}/export?format=md"
    md_dl = requests.get(md_url, headers=headers)
    assert md_dl.status_code == 200
    assert "text/markdown" in md_dl.headers.get("content-type", "")
    assert ".md" in md_dl.headers.get("content-disposition", "")
    assert "Executive Summary & Problem Space" in md_dl.text
    md_path = os.path.join(ARTIFACTS_DIR, "wildfire_report.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md_dl.text)
    print(f"[PASS] 5d. Markdown export verified ({os.path.getsize(md_path)} bytes)")

    # 5E. LaTeX Export
    tex_url = f"{BASE_URL}/api/v1/reports/{report_id}/export?format=latex"
    tex_dl = requests.get(tex_url, headers=headers)
    assert tex_dl.status_code == 200
    assert "latex" in tex_dl.headers.get("content-type", "")
    assert ".tex" in tex_dl.headers.get("content-disposition", "")
    assert "\\documentclass" in tex_dl.text
    assert "\\begin{document}" in tex_dl.text
    tex_path = os.path.join(ARTIFACTS_DIR, "wildfire_report.tex")
    with open(tex_path, "w", encoding="utf-8") as f:
        f.write(tex_dl.text)
    print(f"[PASS] 5e. LaTeX export verified ({os.path.getsize(tex_path)} bytes)")

    # 5F. CSV Evaluation Export
    csv_url = f"{BASE_URL}/api/v1/projects/{project_id}/export/csv"
    csv_dl = requests.get(csv_url, headers=headers)
    assert csv_dl.status_code == 200
    assert "text/csv" in csv_dl.headers.get("content-type", "")
    assert ".csv" in csv_dl.headers.get("content-disposition", "")
    assert csv_dl.content.startswith(b"\xef\xbb\xbf"), "Missing UTF-8 BOM marker for Excel compatibility"
    csv_path = os.path.join(ARTIFACTS_DIR, "evaluation_scores.csv")
    with open(csv_path, "wb") as f:
        f.write(csv_dl.content)
    
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        csv_rows = list(reader)
    assert len(csv_rows) >= len(ideas), f"Expected at least {len(ideas)} rows in CSV, got {len(csv_rows)}"
    first_row = csv_rows[0]
    expected_cols = [
        "Project ID", "Project Title", "Idea ID", "Idea Title",
        "Composite Score", "Grade", "Strategic Recommendation",
        "Problem Significance", "Target User Clarity", "Technical Feasibility",
        "Implementation Complexity", "Estimated Development Cost",
        "Adoption Potential", "Competitive Differentiation", "Novelty Distance",
        "Security Privacy Risks", "Evidence Quality Confidence"
    ]
    for c in expected_cols:
        assert c in first_row, f"Missing CSV column: {c}"
    print(f"[PASS] 5f. CSV export verified ({len(csv_rows)} idea rows, UTF-8-sig BOM, all 10 dimension columns)")

    # 5G. Test Direct Download with ?token= Query Parameter (Browser Link Compatibility)
    print("\n--- Testing Direct Token Query Parameter Download ---")
    query_pdf_dl = requests.get(f"{BASE_URL}/api/v1/reports/{report_id}/export?format=pdf&token={token}")
    assert query_pdf_dl.status_code == 200, f"Token query download failed: {query_pdf_dl.status_code}"
    assert query_pdf_dl.content.startswith(b"%PDF-")
    print("[PASS] 5g. Direct ?token= query parameter download succeeds without Authorization header")

    # -------------------------------------------------------------------------
    # 6. Persistence & Reopening Verification
    # -------------------------------------------------------------------------
    print("\n--- Verifying Database Persistence & Report Reopening ---")
    reps_list = requests.get(f"{BASE_URL}/api/v1/projects/{project_id}/reports", headers=headers).json()
    assert any(r["id"] == report_id for r in reps_list), "Report not listed in project reports"
    
    get_rep = requests.get(f"{BASE_URL}/api/v1/reports/{report_id}", headers=headers).json()
    assert get_rep["id"] == report_id
    assert get_rep["title"] == report["title"]
    assert get_rep["content_markdown"] == report["content_markdown"]
    print("[PASS] 6a. Saved report retrieved from database; title and full Markdown match exactly")

    # -------------------------------------------------------------------------
    # 7. Security & Multi-Tenant Isolation
    # -------------------------------------------------------------------------
    print("\n--- Verifying Multi-Tenant Isolation & Path Privacy ---")
    intruder_email = f"intruder_{uuid.uuid4().hex[:6]}@external.io"
    reg_i = requests.post(f"{BASE_URL}/api/v1/auth/register", json={
        "email": intruder_email,
        "password": "IntruderPassword123!",
        "full_name": "Intruder User"
    })
    assert reg_i.status_code in [200, 201]
    tok_i = requests.post(f"{BASE_URL}/api/v1/auth/login", json={
        "email": intruder_email,
        "password": "IntruderPassword123!"
    }).json()["access_token"]
    intruder_headers = {"Authorization": f"Bearer {tok_i}"}

    # Intruder tries to access user 1's project data and reports
    blocked_endpoints = [
        requests.get(f"{BASE_URL}/api/v1/projects/{project_id}/reports", headers=intruder_headers),
        requests.get(f"{BASE_URL}/api/v1/reports/{report_id}", headers=intruder_headers),
        requests.get(f"{BASE_URL}/api/v1/reports/{report_id}/export?format=pdf", headers=intruder_headers),
        requests.get(f"{BASE_URL}/api/v1/presentations/{pres_id}/export?format=pptx", headers=intruder_headers),
        requests.get(f"{BASE_URL}/api/v1/projects/{project_id}/export/csv", headers=intruder_headers),
        requests.delete(f"{BASE_URL}/api/v1/reports/{report_id}", headers=intruder_headers)
    ]
    for r in blocked_endpoints:
        assert r.status_code in [403, 404], f"Tenant isolation breach! Received HTTP {r.status_code} on {r.url}"
    print("[PASS] 7a. Tenant isolation enforced: unauthorized user received 403/404 on all endpoints")

    # Path Privacy Check: ensure no server internal paths exposed
    for resp in [pdf_dl, docx_dl, pptx_dl, csv_dl]:
        for header, val in resp.headers.items():
            assert "C:\\" not in val and "/Users/" not in val and "/home/" not in val, f"Internal path leaked in header {header}: {val}"
    print("[PASS] 7b. No internal file paths exposed in HTTP headers")

    # -------------------------------------------------------------------------
    # 8. Regression Verification of Phase 4, Phase 5, Phase 6
    # -------------------------------------------------------------------------
    print("\n--- Regression Verification of Previous Phases ---")
    
    # Phase 4: Scoring evaluation
    eval_probe = requests.post(f"{BASE_URL}/api/v1/projects/{project_id}/evaluate", json={
        "idea_ids": [idea_id]
    }, headers=headers)
    assert eval_probe.status_code in [200, 201], f"Phase 4 scoring failed: {eval_probe.text}"
    print("[PASS] 8a. Phase 4: Scoring engine operational")

    # Phase 5: Scholarly Literature Search & Patent Search
    paper_search = requests.get(f"{BASE_URL}/api/v1/research/papers/search?query=wildfire+detection+infrared&limit=3", headers=headers)
    assert paper_search.status_code == 200, f"Paper search failed: {paper_search.text}"
    papers_found = paper_search.json().get("results", [])
    assert len(papers_found) > 0, "Expected at least 1 academic paper from OpenAlex/Crossref"
    
    patent_sources = requests.get(f"{BASE_URL}/api/v1/research/patents/sources", headers=headers)
    assert patent_sources.status_code == 200, f"Patent sources check failed: {patent_sources.text}"
    
    patent_search = requests.get(f"{BASE_URL}/api/v1/research/patents/search?query=autonomous+drone+infrared&limit=3", headers=headers)
    if patent_search.status_code == 503:
        assert "USPTO Open Data Portal is not configured" in patent_search.text or "not configured" in patent_search.text
        print(f"[PASS] 8b. Phase 5: Academic paper search active ({len(papers_found)} papers retrieved from OpenAlex); Patent provider correctly reporting unconfigured API key status")
    else:
        assert patent_search.status_code == 200
        print(f"[PASS] 8b. Phase 5: Academic & Patent search operational ({len(papers_found)} papers, {len(patent_search.json().get('results', []))} patents found)")

    # Phase 6: AI Mentor Chat & Implementation Roadmap
    mentor_res = requests.post(f"{BASE_URL}/api/v1/mentor/conversations", json={
        "project_id": project_id,
        "idea_id": idea_id,
        "mentor_persona": "dr_arun",
        "initial_message": "What are the primary operational challenges in deploying drone mesh networks for wildfire detection?"
    }, headers=headers)
    assert mentor_res.status_code in [200, 201], f"Mentor chat failed: {mentor_res.text}"
    
    roadmap_res = requests.post(f"{BASE_URL}/api/v1/projects/{project_id}/roadmap/generate", json={
        "idea_id": idea_id,
        "target_weeks": 12
    }, headers=headers)
    assert roadmap_res.status_code in [200, 201], f"Roadmap generation failed: {roadmap_res.text}"
    print("[PASS] 8c. Phase 6: AI Mentor conversation and Roadmap generation operational")

    print("\n" + "=" * 70)
    print("ALL END-TO-END VERIFICATION CHECKS COMPLETED WITH 100% SUCCESS!")
    print("=" * 70)

if __name__ == "__main__":
    verify_all()
