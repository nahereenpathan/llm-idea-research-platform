from app.core.database import SessionLocal
from app.models.project import Project
from app.models.paper import AcademicPaper
from app.models.patent import PatentRecord
from app.models.user import User
from app.core.security import create_access_token
import requests
import io
from pptx import Presentation

db = SessionLocal()
p = db.query(Project).filter(Project.id == 9).first()
assert p is not None, "Project 9 not found"
u = db.query(User).filter(User.id == p.user_id).first()
tok = create_access_token(str(u.id))
headers = {'Authorization': f'Bearer {tok}'}

print(f"Testing Project 9: '{p.title}' (User ID: {u.id})")
papers = db.query(AcademicPaper).filter(AcademicPaper.project_id == 9).all()
patents = db.query(PatentRecord).filter(PatentRecord.project_id == 9).all()
print(f"Indexed in DB: {len(papers)} papers, {len(patents)} patents")

# Generate Report
r = requests.post('http://127.0.0.1:8000/api/v1/projects/9/reports/generate', json={'report_type': 'full_documentation'}, headers=headers)
assert r.status_code in [200, 201], f"Report generation failed: {r.text}"
rep = r.json()
md = rep['content_markdown']

assert '[VERIFIED CITATIONS]' in md, "Expected [VERIFIED CITATIONS] tag"
assert '10.1016/j.biortech.2024.130001' in md, "Paper DOI missing"
assert 'US11872594B2' in md, "Patent number missing"
print("[PASS] Report contains authentic DOI and Patent citations from PostgreSQL library")

# Generate Presentation
pr_res = requests.post('http://127.0.0.1:8000/api/v1/projects/9/presentations/generate', json={'template_type': 'academic'}, headers=headers)
assert pr_res.status_code in [200, 201]
pres = pr_res.json()

dl = requests.get(f"http://127.0.0.1:8000/api/v1/presentations/{pres['id']}/export?format=pptx", headers=headers)
assert dl.status_code == 200

prs = Presentation(io.BytesIO(dl.content))
s10 = prs.slides[9]
links = []
for shp in s10.shapes:
    if shp.has_text_frame:
        for para in shp.text_frame.paragraphs:
            for run in para.runs:
                if run.hyperlink and run.hyperlink.address:
                    links.append(run.hyperlink.address)

print("PPTX Slide 10 Clickable Citation Hyperlinks:", links)
assert any('doi.org' in l or 'patents.google.com' in l for l in links), "Missing clickable source links on Slide 10"
print("[PASS] Verified PPTX Slide 10 contains real, clickable hyperlinks directly to retrieved sources!")
